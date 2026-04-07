"""
=== BAI PHONG VAN GENAI - CODE TREN MAY ===

De bai:
  User gui file "report.docx" va yeu cau:
  "Trich xuat tat ca bang doanh thu trong file nay, tao file Excel, roi gui qua email cho sep toi"

Yeu cau:
  1. Thiet ke va code cac tool can thiet cho agent
  2. Xay dung agent loop de agent tu dong xu ly yeu cau tren
  3. Agent phai biet hoi lai user khi thieu thong tin
  4. Su dung OpenAI API (function calling)

Tieu chi danh gia:
  - Thiet ke tool hop ly (du tool, khong thua)          [3 diem]
  - Agent loop dung (Thought -> Action -> Observation)   [3 diem]
  - Xu ly thieu thong tin (hoi lai user)                 [2 diem]
  - Code chay duoc, cho ket qua dung                     [2 diem]

Luu y:
  - File report.docx da duoc cung cap san
  - Ham send_email chi can mo phong bang print, KHONG can gui email that
  - OPENAI_API_KEY da duoc set san trong moi truong
  - Thoi gian: 15 phut

Cai dat thu vien:
  pip install openai python-docx openpyxl

Chay: python starter.py
"""

import os
import json
import docx
import openpyxl
from openai import OpenAI

# API KEY mặc định (đã cung cấp trong đề bài)
os.environ["OPENAI_API_KEY"] = os.getenv("OPENAI_API_KEY")
client = OpenAI()

# ==========================================
# 1. THIẾT KẾ & CODE CÁC TOOL CẦN THIẾT
# ==========================================

def extract_tables_from_docx(file_path: str) -> str:
    """Đọc file docx và trích xuất tất cả các bảng thành chuỗi JSON."""
    try:
        doc = docx.Document(file_path)
        tables_data = []
        for i, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            tables_data.append({"table_name": f"Bang_{i+1}", "data": table_data})
        return json.dumps(tables_data, ensure_ascii=False)
    except Exception as e:
        return f"Error reading document: {str(e)}. Hãy kiểm tra xem file {file_path} có tồn tại không."

def create_excel_file(json_data: str, output_file: str) -> str:
    """Tạo file Excel từ dữ liệu JSON (được trích xuất từ Word)."""
    try:
        data = json.loads(json_data)
        wb = openpyxl.Workbook()
        wb.remove(wb.active) # Xóa sheet mặc định

        for table in data:
            ws = wb.create_sheet(title=table.get("table_name", "Sheet"))
            for row in table.get("data", []):
                ws.append(row)

        wb.save(output_file)
        return f"Success: Đã tạo file Excel thành công tại {output_file}"
    except Exception as e:
        return f"Error creating Excel: {str(e)}"

def send_email(to_email: str, subject: str, body: str, attachment_file: str) -> str:
    """Mô phỏng hàm gửi email bằng print."""
    print("\n" + "="*45)
    print("📧 [MOCK EMAIL SENT]")
    print(f"To:         {to_email}")
    print(f"Subject:    {subject}")
    print(f"Attachment: {attachment_file}")
    print(f"Body:       {body}")
    print("="*45 + "\n")
    return f"Success: Email đã được gửi tới {to_email}"


# ==========================================
# 2. KHAI BÁO TOOLS CHO OPENAI (FUNCTION CALLING)
# ==========================================
tools = [
    {
        "type": "function",
        "function": {
            "name": "extract_tables_from_docx",
            "description": "Trích xuất dữ liệu bảng từ file Word (.docx). Trả về dữ liệu dạng chuỗi JSON.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Đường dẫn tới file Word (ví dụ: report.docx)"
                    }
                },
                "required": ["file_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "create_excel_file",
            "description": "Tạo file Excel (.xlsx) từ dữ liệu JSON chứa cấu trúc bảng.",
            "parameters": {
                "type": "object",
                "properties": {
                    "json_data": {
                        "type": "string",
                        "description": "Chuỗi JSON chứa dữ liệu bảng cần ghi vào Excel."
                    },
                    "output_file": {
                        "type": "string",
                        "description": "Tên file Excel đầu ra (ví dụ: doanh_thu.xlsx)"
                    }
                },
                "required": ["json_data", "output_file"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "send_email",
            "description": "Gửi email đính kèm file cho một người nào đó.",
            "parameters": {
                "type": "object",
                "properties": {
                    "to_email": {
                        "type": "string",
                        "description": "Địa chỉ email người nhận. (Bắt buộc phải có định dạng email hợp lệ)."
                    },
                    "subject": {
                        "type": "string",
                        "description": "Tiêu đề email."
                    },
                    "body": {
                        "type": "string",
                        "description": "Nội dung email."
                    },
                    "attachment_file": {
                        "type": "string",
                        "description": "Tên file đính kèm để gửi đi."
                    }
                },
                "required": ["to_email", "subject", "body", "attachment_file"]
            }
        }
    }
]


# ==========================================
# 3. XÂY DỰNG AGENT LOOP
# ==========================================
def run_agent(user_prompt: str):
    """
    Vòng lặp Agent thực hiện: Thought -> Action -> Observation
    Và hỏi lại User khi thiếu thông tin (ví dụ: thiếu email).
    """
    messages = [
        {
            "role": "system",
            "content": (
                "Bạn là một trợ lý AI thông minh chuyên tự động hóa công việc văn phòng. "
                "Quy trình bạn cần làm:\n"
                "1. Trích xuất bảng từ file Word (.docx).\n"
                "2. Tạo file Excel từ dữ liệu vừa trích xuất.\n"
                "3. Gửi file Excel đó qua email cho người được chỉ định.\n\n"
                "QUAN TRỌNG (Xử lý thiếu thông tin): Nếu người dùng yêu cầu gửi email cho 'sếp' "
                "hoặc bất kỳ ai mà chưa cung cấp địa chỉ email cụ thể, BẠN PHẢI HỎI LẠI NGƯỜI DÙNG "
                "địa chỉ email là gì trước khi gọi hàm send_email. Tuyệt đối không tự bịa ra email."
            )
        },
        {"role": "user", "content": user_prompt}
    ]

    while True:
        # THOUGHT: Gọi API để LLM suy nghĩ và đưa ra quyết định
        response = client.chat.completions.create(
            model="gpt-4o", # Có thể thay bằng gpt-3.5-turbo nếu cần tối ưu chi phí
            messages=messages,
            tools=tools,
            tool_choice="auto",
        )
        
        response_message = response.choices[0].message

        # ACTION: Nếu Agent quyết định gọi tool (function calling)
        if response_message.tool_calls:
            messages.append(response_message) # Lưu lại quyết định của agent vào lịch sử
            
            for tool_call in response_message.tool_calls:
                function_name = tool_call.function.name
                function_args = json.loads(tool_call.function.arguments)

                print(f"🧠 [Thought/Action] Agent quyết định gọi hàm: '{function_name}' với tham số: {function_args}")

                # Thực thi hàm tương ứng
                if function_name == "extract_tables_from_docx":
                    tool_result = extract_tables_from_docx(function_args.get("file_path"))
                elif function_name == "create_excel_file":
                    tool_result = create_excel_file(function_args.get("json_data"), function_args.get("output_file"))
                elif function_name == "send_email":
                    tool_result = send_email(
                        to_email=function_args.get("to_email"),
                        subject=function_args.get("subject"),
                        body=function_args.get("body"),
                        attachment_file=function_args.get("attachment_file")
                    )
                else:
                    tool_result = f"Error: Hàm {function_name} không tồn tại."

                # OBSERVATION: Trả kết quả thực thi về cho Agent
                print(f"👀 [Observation] Kết quả từ tool: {str(tool_result)[:150]}...\n")
                messages.append({
                    "tool_call_id": tool_call.id,
                    "role": "tool",
                    "name": function_name,
                    "content": str(tool_result),
                })
                
        # GIAO TIẾP VỚI USER: Agent không gọi tool mà trực tiếp trả lời (Hỏi thông tin hoặc Báo cáo xong)
        else:
            agent_response = response_message.content
            print(f"🤖 Agent: {agent_response}")

            # Đợi input từ user. Nếu Agent đã hoàn thành (không có câu hỏi), user có thể gõ exit để thoát.
            user_input = input("\n👤 User (Nhập câu trả lời hoặc 'exit' để dừng): ")
            if user_input.lower().strip() in ['exit', 'quit', '']:
                print("Đã kết thúc chương trình.")
                break
            
            # Cập nhật context bằng câu trả lời của user để tiếp tục vòng lặp
            messages.append({"role": "user", "content": user_input})


if __name__ == "__main__":
    user_request = (
        'Trich xuat tat ca bang doanh thu trong file "report.docx", '
        'tao file Excel, roi gui qua email cho sep toi.'
    )
    print(f"👤 User: {user_request}\n")
    
    # Kích hoạt agent loop
    run_agent(user_request)